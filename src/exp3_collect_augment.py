from __future__ import annotations

import csv
import io
import json
import random
import shutil
import urllib.request
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image, ImageDraw, ImageEnhance, ImageFont, ImageOps


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DATASET_DIR = PROJECT_ROOT / "dataset" / "exp3"
DOWNLOAD_DIR = DATASET_DIR / "_downloads"
RAW_DIR = DATASET_DIR / "raw"
AUGMENTED_DIR = DATASET_DIR / "augmented"
OUTPUT_DIR = PROJECT_ROOT / "outputs" / "exp3"

ZENODO_RECORD_API = "https://zenodo.org/api/records/3271625"
ZENODO_RECORD_PAGE = "https://zenodo.org/records/3271625"
ZENODO_DOI = "10.5281/zenodo.3271625"
REPO_URL = "https://github.com/Lovertin/cv_exp.git"
SOURCE_FILE = "src/exp3_collect_augment.py"
LOCAL_SOURCE_PATH = PROJECT_ROOT / SOURCE_FILE

IMAGE_SIZE = 512
RAW_PER_CLASS = 20
AUGMENTATIONS_PER_RAW = 4
RANDOM_SEED = 20260604

CLASS_ZIPS = {
    "open_hand": "raw_open_hand.zip",
    "fist": "raw_fist.zip",
    "right_hand": "raw_right_hand.zip",
}

IMAGE_SUFFIXES = {".jpg", ".jpeg", ".png", ".bmp", ".tif", ".tiff"}


def reset_dir(path: Path) -> None:
    resolved = path.resolve()
    dataset_root = DATASET_DIR.resolve()
    output_root = OUTPUT_DIR.resolve()
    if dataset_root not in resolved.parents and resolved != dataset_root:
        if output_root not in resolved.parents and resolved != output_root:
            raise ValueError(f"Refusing to reset unexpected path: {resolved}")
    if path.exists():
        shutil.rmtree(path)
    path.mkdir(parents=True, exist_ok=True)


def download_zenodo_files() -> dict[str, dict[str, object]]:
    DOWNLOAD_DIR.mkdir(parents=True, exist_ok=True)
    with urllib.request.urlopen(ZENODO_RECORD_API, timeout=60) as response:
        record = json.load(response)

    files = {item["key"]: item for item in record["files"]}
    downloaded: dict[str, dict[str, object]] = {}
    for class_name, zip_name in CLASS_ZIPS.items():
        if zip_name not in files:
            raise FileNotFoundError(f"{zip_name} not found in Zenodo record")
        file_info = files[zip_name]
        url = file_info["links"]["self"]
        size = int(file_info["size"])
        zip_path = DOWNLOAD_DIR / zip_name
        if not zip_path.exists() or zip_path.stat().st_size != size:
            print(f"Downloading {zip_name} ...")
            with urllib.request.urlopen(url, timeout=120) as response:
                zip_path.write_bytes(response.read())
        downloaded[class_name] = {
            "zip_name": zip_name,
            "path": str(zip_path),
            "size": size,
            "url": url,
        }
    return downloaded


def load_image_from_zip(zip_file: ZipFile, member_name: str) -> Image.Image:
    with zip_file.open(member_name) as fp:
        data = fp.read()
    image = Image.open(io.BytesIO(data))
    image = ImageOps.exif_transpose(image)
    return image.convert("RGB")


def normalize_image(image: Image.Image, size: int = IMAGE_SIZE) -> Image.Image:
    image = ImageOps.contain(image, (size, size), Image.Resampling.LANCZOS)
    canvas = Image.new("RGB", (size, size), (24, 24, 24))
    x = (size - image.width) // 2
    y = (size - image.height) // 2
    canvas.paste(image, (x, y))
    return canvas


def select_zip_members(zip_path: Path, count: int, class_name: str) -> list[str]:
    with ZipFile(zip_path) as zip_file:
        members = [
            name
            for name in zip_file.namelist()
            if Path(name).suffix.lower() in IMAGE_SUFFIXES and not name.endswith("/")
        ]
    if len(members) < count:
        raise ValueError(f"{zip_path.name} has only {len(members)} images")
    rng = random.Random(RANDOM_SEED + sum(ord(ch) for ch in class_name))
    return sorted(rng.sample(sorted(members), count))


def collect_raw_images(downloaded: dict[str, dict[str, object]]) -> list[dict[str, str]]:
    reset_dir(RAW_DIR)
    manifest: list[dict[str, str]] = []
    for class_name, info in downloaded.items():
        class_dir = RAW_DIR / class_name
        class_dir.mkdir(parents=True, exist_ok=True)
        zip_path = Path(str(info["path"]))
        selected_members = select_zip_members(zip_path, RAW_PER_CLASS, class_name)
        with ZipFile(zip_path) as zip_file:
            for index, member_name in enumerate(selected_members, start=1):
                image = load_image_from_zip(zip_file, member_name)
                normalized = normalize_image(image)
                output_name = f"{class_name}_{index:03d}.jpg"
                output_path = class_dir / output_name
                normalized.save(output_path, quality=92)
                manifest.append(
                    {
                        "class": class_name,
                        "raw_file": str(output_path.relative_to(PROJECT_ROOT)),
                        "source_zip": str(info["zip_name"]),
                        "source_member": member_name,
                    }
                )
    return manifest


def random_scale_and_translate(image: Image.Image, rng: random.Random) -> Image.Image:
    width, height = image.size
    scale = rng.uniform(0.82, 1.18)
    scaled = image.resize(
        (max(1, int(width * scale)), max(1, int(height * scale))),
        Image.Resampling.BICUBIC,
    )
    canvas = Image.new("RGB", (width, height), (24, 24, 24))
    max_dx = int(width * 0.08)
    max_dy = int(height * 0.08)
    x = (width - scaled.width) // 2 + rng.randint(-max_dx, max_dx)
    y = (height - scaled.height) // 2 + rng.randint(-max_dy, max_dy)
    canvas.paste(scaled, (x, y))
    if scaled.width > width or scaled.height > height:
        left = max(0, -x)
        top = max(0, -y)
        right = min(scaled.width, left + width)
        bottom = min(scaled.height, top + height)
        cropped = scaled.crop((left, top, right, bottom))
        canvas = Image.new("RGB", (width, height), (24, 24, 24))
        canvas.paste(cropped, (max(0, x), max(0, y)))
    return canvas


def crop_random_region(image: Image.Image, rng: random.Random) -> Image.Image:
    width, height = image.size
    crop_ratio = rng.uniform(0.88, 1.0)
    crop_w = int(width * crop_ratio)
    crop_h = int(height * crop_ratio)
    left = rng.randint(0, width - crop_w)
    top = rng.randint(0, height - crop_h)
    cropped = image.crop((left, top, left + crop_w, top + crop_h))
    return cropped.resize((width, height), Image.Resampling.BICUBIC)


def color_jitter(image: Image.Image, rng: random.Random) -> Image.Image:
    image = ImageEnhance.Brightness(image).enhance(rng.uniform(0.78, 1.24))
    image = ImageEnhance.Contrast(image).enhance(rng.uniform(0.78, 1.28))
    image = ImageEnhance.Color(image).enhance(rng.uniform(0.70, 1.35))
    return image


def add_light_noise(image: Image.Image, rng: random.Random) -> Image.Image:
    pixels = image.load()
    width, height = image.size
    noise_strength = rng.randint(3, 12)
    stride = rng.randint(3, 6)
    for y in range(0, height, stride):
        for x in range(0, width, stride):
            r, g, b = pixels[x, y]
            delta = rng.randint(-noise_strength, noise_strength)
            pixels[x, y] = (
                max(0, min(255, r + delta)),
                max(0, min(255, g + delta)),
                max(0, min(255, b + delta)),
            )
    return image


def augment_image(image: Image.Image, rng: random.Random) -> Image.Image:
    augmented = image.copy()
    if rng.random() < 0.5:
        augmented = ImageOps.mirror(augmented)
    augmented = crop_random_region(augmented, rng)
    augmented = random_scale_and_translate(augmented, rng)
    angle = rng.uniform(-14.0, 14.0)
    augmented = augmented.rotate(
        angle,
        resample=Image.Resampling.BICUBIC,
        expand=False,
        fillcolor=(24, 24, 24),
    )
    augmented = color_jitter(augmented, rng)
    if rng.random() < 0.55:
        augmented = add_light_noise(augmented, rng)
    return augmented


def augment_dataset(raw_manifest: list[dict[str, str]]) -> list[dict[str, str]]:
    reset_dir(AUGMENTED_DIR)
    manifest: list[dict[str, str]] = []
    for raw_index, item in enumerate(raw_manifest, start=1):
        class_name = item["class"]
        raw_path = PROJECT_ROOT / item["raw_file"]
        output_class_dir = AUGMENTED_DIR / class_name
        output_class_dir.mkdir(parents=True, exist_ok=True)
        image = Image.open(raw_path).convert("RGB")
        for aug_index in range(1, AUGMENTATIONS_PER_RAW + 1):
            rng = random.Random(RANDOM_SEED + raw_index * 101 + aug_index * 17)
            augmented = augment_image(image, rng)
            output_name = f"{raw_path.stem}_aug_{aug_index:02d}.jpg"
            output_path = output_class_dir / output_name
            augmented.save(output_path, quality=92)
            manifest.append(
                {
                    "class": class_name,
                    "augmented_file": str(output_path.relative_to(PROJECT_ROOT)),
                    "source_raw_file": item["raw_file"],
                    "operations": "random crop, random scale, translation, optional flip, rotation, color jitter, light noise",
                }
            )
    return manifest


def write_csv(path: Path, rows: list[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if not rows:
        return
    with path.open("w", encoding="utf-8-sig", newline="") as fp:
        writer = csv.DictWriter(fp, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)


def count_by_class(rows: list[dict[str, str]]) -> dict[str, int]:
    counts: dict[str, int] = {class_name: 0 for class_name in CLASS_ZIPS}
    for row in rows:
        counts[row["class"]] = counts.get(row["class"], 0) + 1
    return counts


def get_font(size: int) -> ImageFont.ImageFont:
    for font_name in ["arial.ttf", "calibri.ttf"]:
        try:
            return ImageFont.truetype(font_name, size)
        except OSError:
            continue
    return ImageFont.load_default()


def make_contact_sheet(
    image_paths: list[Path],
    labels: list[str],
    output_path: Path,
    columns: int,
    thumb_size: tuple[int, int] = (170, 170),
) -> None:
    label_height = 34
    rows = (len(image_paths) + columns - 1) // columns
    width = columns * thumb_size[0]
    height = rows * (thumb_size[1] + label_height)
    sheet = Image.new("RGB", (width, height), (245, 245, 245))
    draw = ImageDraw.Draw(sheet)
    font = get_font(14)
    for index, path in enumerate(image_paths):
        row, column = divmod(index, columns)
        x = column * thumb_size[0]
        y = row * (thumb_size[1] + label_height)
        with Image.open(path) as image:
            image = image.convert("RGB")
            image.thumbnail((thumb_size[0] - 12, thumb_size[1] - 12))
            paste_x = x + (thumb_size[0] - image.width) // 2
            paste_y = y + (thumb_size[1] - image.height) // 2
            sheet.paste(image, (paste_x, paste_y))
        draw.text((x + 6, y + thumb_size[1] + 6), labels[index], fill=(0, 0, 0), font=font)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    sheet.save(output_path, quality=92)


def make_counts_chart(raw_counts: dict[str, int], aug_counts: dict[str, int], output_path: Path) -> None:
    width, height = 760, 420
    margin_left, margin_bottom, margin_top = 90, 70, 60
    chart_w = width - margin_left - 40
    chart_h = height - margin_top - margin_bottom
    image = Image.new("RGB", (width, height), (255, 255, 255))
    draw = ImageDraw.Draw(image)
    title_font = get_font(24)
    label_font = get_font(16)
    small_font = get_font(14)
    draw.text((margin_left, 18), "Experiment 3 Dataset Counts", fill=(0, 0, 0), font=title_font)
    max_value = max(max(raw_counts.values()), max(aug_counts.values()))
    scale = chart_h / max_value
    classes = list(CLASS_ZIPS.keys())
    group_w = chart_w / len(classes)
    bar_w = 42
    colors = {"raw": (70, 130, 180), "aug": (46, 160, 110)}
    y0 = margin_top + chart_h
    draw.line((margin_left, margin_top, margin_left, y0), fill=(30, 30, 30), width=2)
    draw.line((margin_left, y0, width - 30, y0), fill=(30, 30, 30), width=2)
    for tick in range(0, max_value + 1, 20):
        y = y0 - tick * scale
        draw.line((margin_left - 5, y, margin_left, y), fill=(30, 30, 30), width=1)
        draw.text((20, y - 8), str(tick), fill=(30, 30, 30), font=small_font)
    for idx, class_name in enumerate(classes):
        group_x = margin_left + idx * group_w + group_w / 2
        raw_value = raw_counts[class_name]
        aug_value = aug_counts[class_name]
        for offset, value, color, kind in [
            (-bar_w - 4, raw_value, colors["raw"], "raw"),
            (4, aug_value, colors["aug"], "aug"),
        ]:
            x1 = int(group_x + offset)
            x2 = int(x1 + bar_w)
            y1 = int(y0 - value * scale)
            draw.rectangle((x1, y1, x2, y0), fill=color)
            draw.text((x1, y1 - 18), str(value), fill=(0, 0, 0), font=small_font)
            draw.text((x1 + 4, y0 + 8), kind, fill=(0, 0, 0), font=small_font)
        draw.text((int(group_x - 48), y0 + 30), class_name, fill=(0, 0, 0), font=label_font)
    image.save(output_path, quality=92)


def build_visual_outputs(raw_manifest: list[dict[str, str]], aug_manifest: list[dict[str, str]]) -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    raw_samples: list[Path] = []
    raw_labels: list[str] = []
    for class_name in CLASS_ZIPS:
        class_items = [row for row in raw_manifest if row["class"] == class_name][:5]
        raw_samples.extend(PROJECT_ROOT / row["raw_file"] for row in class_items)
        raw_labels.extend([class_name] * len(class_items))
    make_contact_sheet(
        raw_samples,
        raw_labels,
        OUTPUT_DIR / "01_raw_samples_summary.jpg",
        columns=5,
    )

    aug_samples: list[Path] = []
    aug_labels: list[str] = []
    for class_name in CLASS_ZIPS:
        first_raw = next(row for row in raw_manifest if row["class"] == class_name)
        aug_samples.append(PROJECT_ROOT / first_raw["raw_file"])
        aug_labels.append(f"{class_name} raw")
        related_aug = [row for row in aug_manifest if row["source_raw_file"] == first_raw["raw_file"]][:4]
        aug_samples.extend(PROJECT_ROOT / row["augmented_file"] for row in related_aug)
        aug_labels.extend([f"{class_name} aug{i}" for i in range(1, len(related_aug) + 1)])
    make_contact_sheet(
        aug_samples,
        aug_labels,
        OUTPUT_DIR / "02_augmentation_summary.jpg",
        columns=5,
    )

    make_counts_chart(
        count_by_class(raw_manifest),
        count_by_class(aug_manifest),
        OUTPUT_DIR / "03_dataset_counts_summary.jpg",
    )


def write_summary(
    downloaded: dict[str, dict[str, object]],
    raw_manifest: list[dict[str, str]],
    aug_manifest: list[dict[str, str]],
) -> None:
    raw_counts = count_by_class(raw_manifest)
    aug_counts = count_by_class(aug_manifest)
    summary = {
        "source": {
            "name": "Hand gestures raw images",
            "record_page": ZENODO_RECORD_PAGE,
            "doi": ZENODO_DOI,
            "api": ZENODO_RECORD_API,
            "files": downloaded,
        },
        "parameters": {
            "classes": list(CLASS_ZIPS.keys()),
            "raw_per_class": RAW_PER_CLASS,
            "augmentations_per_raw": AUGMENTATIONS_PER_RAW,
            "image_size": IMAGE_SIZE,
            "random_seed": RANDOM_SEED,
        },
        "counts": {
            "raw_by_class": raw_counts,
            "augmented_by_class": aug_counts,
            "raw_total": len(raw_manifest),
            "augmented_total": len(aug_manifest),
        },
        "outputs": [
            "01_raw_samples_summary.jpg",
            "02_augmentation_summary.jpg",
            "03_dataset_counts_summary.jpg",
            "raw_manifest.csv",
            "augmentation_manifest.csv",
            "exp3_result_summary.txt",
            "exp3_report_content.docx",
            "实验三报告内容.md",
        ],
    }
    (OUTPUT_DIR / "exp3_summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    lines = [
        "Experiment 3: Data Collection and Augmentation",
        "",
        "Data source: Hand gestures raw images",
        f"Zenodo page: {ZENODO_RECORD_PAGE}",
        f"DOI: {ZENODO_DOI}",
        "",
        "Selected classes:",
    ]
    for class_name in CLASS_ZIPS:
        lines.append(
            f"- {class_name}: raw={raw_counts[class_name]}, augmented={aug_counts[class_name]}"
        )
    lines.extend(
        [
            "",
            f"Raw image total: {len(raw_manifest)}",
            f"Augmented image total: {len(aug_manifest)}",
            f"Image size: {IMAGE_SIZE}x{IMAGE_SIZE}",
            f"Random seed: {RANDOM_SEED}",
            "",
            "Generated files:",
        ]
    )
    for path in sorted(OUTPUT_DIR.glob("*")):
        if path.is_file():
            lines.append(f"- {path.name}")
    (OUTPUT_DIR / "exp3_result_summary.txt").write_text(
        "\n".join(lines), encoding="utf-8"
    )


def make_report_markdown(raw_total: int, aug_total: int, raw_counts: dict[str, int], aug_counts: dict[str, int]) -> str:
    class_lines = "\n".join(
        f"- `{class_name}`：原始 {raw_counts[class_name]} 张，增强 {aug_counts[class_name]} 张"
        for class_name in CLASS_ZIPS
    )
    return f"""# 实验三  数据的采集与增强

## 一、实验目的

1. 了解图像数据采集的基本流程。
2. 熟悉公开数据集获取和数据整理方法。
3. 掌握随机缩放、翻转、裁剪、平移、旋转和色彩抖动等常见数据增强操作。
4. 为后续数据集制作和目标检测实验准备图像数据。

## 二、实验原理

数据采集是计算机视觉任务的基础。实验中可以通过爬虫软件、公开数据集或人工拍摄等方式获取原始图像。采集后的图像需要进行类别整理、尺寸统一、格式转换和质量检查，保证数据可以被后续程序稳定读取。

数据增强是从已有样本生成更多训练样本的方法。常用增强方式包括随机缩放、水平翻转、随机裁剪、平移、旋转、亮度和对比度扰动、色彩扰动以及轻微噪声。合理的数据增强可以提升样本多样性，降低模型对单一背景、尺度和光照条件的依赖，从而提高模型泛化能力。

## 三、实验结果分析

本实验选择“手势动作”作为数据主题，使用公开数据集 Hand gestures raw images 作为采集来源。该数据集页面为 `{ZENODO_RECORD_PAGE}`，DOI 为 `{ZENODO_DOI}`。实验选择 `open_hand`、`fist`、`right_hand` 三类图像，每类抽取 20 张，共得到 {raw_total} 张原始图像，满足不少于 50 张原始数据的要求。

各类别数量如下：

{class_lines}

增强阶段对每张原始图像生成 {AUGMENTATIONS_PER_RAW} 张增强图像，增强方法包括随机裁剪、随机缩放、水平/垂直方向平移、随机水平翻转、轻微旋转、亮度/对比度/色彩扰动和轻微噪声。最终得到 {aug_total} 张增强图像，超过实验要求的 200 张。

从结果图可以看出，增强后的图像在位置、尺度、角度和光照颜色上与原始图像存在差异，但仍保持了手势类别语义不变。这样的增强数据可用于后续实验四的数据集标注与划分，也可以继续用于实验五目标检测模型的训练。

## 四、程序源代码

源码仓库：

```text
{REPO_URL}
```

本实验源码文件：

```text
{SOURCE_FILE}
```

本地源码路径：

```text
{LOCAL_SOURCE_PATH}
```

## 五、程序说明

运行环境：

```bash
pip install -r requirements.txt
```

运行命令：

```bash
python E:\\实验报告\\实验\\src\\exp3_collect_augment.py
```

程序运行后，原始数据保存到：

```text
E:\\实验报告\\实验\\dataset\\exp3\\raw
```

增强数据保存到：

```text
E:\\实验报告\\实验\\dataset\\exp3\\augmented
```

结果图、统计表和报告内容保存到：

```text
E:\\实验报告\\实验\\outputs\\exp3
```

## 六、结果文件说明

- `01_raw_samples_summary.jpg`：原始采集样本展示图。
- `02_augmentation_summary.jpg`：原图与增强图像对比。
- `03_dataset_counts_summary.jpg`：原始数据和增强数据数量统计图。
- `raw_manifest.csv`：原始图像来源记录。
- `augmentation_manifest.csv`：增强图像与原始图像对应关系。
- `exp3_result_summary.txt`：实验参数和输出文件摘要。
"""


def set_run_font(run, size: float = 10.5, bold: bool = False) -> None:
    run.font.name = "SimSun"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    run.bold = bold


def add_para(document: Document, text: str, size: float = 10.5):
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    set_run_font(run, size=size)
    return paragraph


def add_heading(document: Document, text: str, level: int = 1):
    paragraph = document.add_heading("", level=level)
    run = paragraph.add_run(text)
    set_run_font(run, size=14 if level == 1 else 12, bold=True)
    return paragraph


def add_code_block(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    run.font.size = Pt(9.5)


def build_report_docx(raw_total: int, aug_total: int, raw_counts: dict[str, int], aug_counts: dict[str, int]) -> Document:
    document = Document()
    document.styles["Normal"].font.name = "SimSun"
    document.styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    document.styles["Normal"].font.size = Pt(10.5)

    add_heading(document, "实验三  数据的采集与增强", 1)

    add_heading(document, "一、实验目的", 2)
    add_para(document, "1. 了解图像数据采集的基本流程。")
    add_para(document, "2. 熟悉公开数据集获取和数据整理方法。")
    add_para(document, "3. 掌握随机缩放、翻转、裁剪、平移、旋转和色彩抖动等常见数据增强操作。")
    add_para(document, "4. 为后续数据集制作和目标检测实验准备图像数据。")

    add_heading(document, "二、实验原理", 2)
    add_para(
        document,
        "数据采集是计算机视觉任务的基础。实验中可以通过爬虫软件、公开数据集或人工拍摄等方式获取原始图像。采集后的图像需要进行类别整理、尺寸统一、格式转换和质量检查，保证数据可以被后续程序稳定读取。",
    )
    add_para(
        document,
        "数据增强是从已有样本生成更多训练样本的方法。常用增强方式包括随机缩放、水平翻转、随机裁剪、平移、旋转、亮度和对比度扰动、色彩扰动以及轻微噪声。合理的数据增强可以提升样本多样性，降低模型对单一背景、尺度和光照条件的依赖，从而提高模型泛化能力。",
    )

    add_heading(document, "三、实验结果分析", 2)
    add_para(
        document,
        f"本实验选择“手势动作”作为数据主题，使用公开数据集 Hand gestures raw images 作为采集来源。数据集页面为 {ZENODO_RECORD_PAGE}，DOI 为 {ZENODO_DOI}。实验选择 open_hand、fist、right_hand 三类图像，每类抽取 20 张，共得到 {raw_total} 张原始图像，满足不少于 50 张原始数据的要求。",
    )
    for class_name in CLASS_ZIPS:
        add_para(
            document,
            f"{class_name}：原始 {raw_counts[class_name]} 张，增强 {aug_counts[class_name]} 张。",
        )
    add_para(
        document,
        f"增强阶段对每张原始图像生成 {AUGMENTATIONS_PER_RAW} 张增强图像，增强方法包括随机裁剪、随机缩放、水平/垂直方向平移、随机水平翻转、轻微旋转、亮度/对比度/色彩扰动和轻微噪声。最终得到 {aug_total} 张增强图像，超过实验要求的 200 张。",
    )
    add_para(
        document,
        "从结果图可以看出，增强后的图像在位置、尺度、角度和光照颜色上与原始图像存在差异，但仍保持了手势类别语义不变。这样的增强数据可用于后续实验四的数据集标注与划分，也可以继续用于实验五目标检测模型的训练。",
    )

    for title, filename in [
        ("图 1 原始采集样本展示", "01_raw_samples_summary.jpg"),
        ("图 2 原图与增强图像对比", "02_augmentation_summary.jpg"),
        ("图 3 数据集数量统计", "03_dataset_counts_summary.jpg"),
    ]:
        add_para(document, title)
        document.add_picture(str(OUTPUT_DIR / filename), width=Inches(6.4))

    add_heading(document, "四、程序源代码", 2)
    add_para(document, "源码仓库：")
    add_code_block(document, REPO_URL)
    add_para(document, "本实验源码文件：")
    add_code_block(document, SOURCE_FILE)
    add_para(document, "本地源码路径：")
    add_code_block(document, str(LOCAL_SOURCE_PATH))

    add_heading(document, "五、程序说明", 2)
    add_para(document, "运行环境：")
    add_code_block(document, "pip install -r requirements.txt")
    add_para(document, "运行命令：")
    add_code_block(document, r"python E:\实验报告\实验\src\exp3_collect_augment.py")
    add_para(document, "程序运行后，原始数据保存到：")
    add_code_block(document, r"E:\实验报告\实验\dataset\exp3\raw")
    add_para(document, "增强数据保存到：")
    add_code_block(document, r"E:\实验报告\实验\dataset\exp3\augmented")
    add_para(document, "结果图、统计表和报告内容保存到：")
    add_code_block(document, r"E:\实验报告\实验\outputs\exp3")

    add_heading(document, "六、结果文件说明", 2)
    add_para(document, "01_raw_samples_summary.jpg：原始采集样本展示图。")
    add_para(document, "02_augmentation_summary.jpg：原图与增强图像对比。")
    add_para(document, "03_dataset_counts_summary.jpg：原始数据和增强数据数量统计图。")
    add_para(document, "raw_manifest.csv：原始图像来源记录。")
    add_para(document, "augmentation_manifest.csv：增强图像与原始图像对应关系。")
    add_para(document, "exp3_result_summary.txt：实验参数和输出文件摘要。")

    return document


def write_report_files(raw_manifest: list[dict[str, str]], aug_manifest: list[dict[str, str]]) -> None:
    raw_counts = count_by_class(raw_manifest)
    aug_counts = count_by_class(aug_manifest)
    raw_total = len(raw_manifest)
    aug_total = len(aug_manifest)
    markdown = make_report_markdown(raw_total, aug_total, raw_counts, aug_counts)
    (OUTPUT_DIR / "实验三报告内容.md").write_text(markdown, encoding="utf-8")
    for name in ["exp3_report_content.docx", "exp3_report_content_fixed.docx"]:
        build_report_docx(raw_total, aug_total, raw_counts, aug_counts).save(OUTPUT_DIR / name)


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    downloaded = download_zenodo_files()
    raw_manifest = collect_raw_images(downloaded)
    aug_manifest = augment_dataset(raw_manifest)
    write_csv(OUTPUT_DIR / "raw_manifest.csv", raw_manifest)
    write_csv(OUTPUT_DIR / "augmentation_manifest.csv", aug_manifest)
    build_visual_outputs(raw_manifest, aug_manifest)
    write_summary(downloaded, raw_manifest, aug_manifest)
    write_report_files(raw_manifest, aug_manifest)
    print((OUTPUT_DIR / "exp3_result_summary.txt").read_text(encoding="utf-8"))


if __name__ == "__main__":
    main()

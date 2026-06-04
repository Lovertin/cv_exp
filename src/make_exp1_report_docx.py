from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = PROJECT_ROOT / "outputs" / "exp1"
REPO_URL = "https://github.com/Lovertin/cv_exp.git"
SOURCE_FILE = "src/exp1_basic_operations.py"
LOCAL_SOURCE_PATH = PROJECT_ROOT / SOURCE_FILE


REPORT_MARKDOWN = f"""# 实验一  基本图像操作

## 一、实验目的

1. 掌握图像处理相关软件和 Python OpenCV 运行环境的基本使用方法。
2. 掌握使用 OpenCV 实现图像读取、写入、显示和缩放操作。
3. 掌握使用 NumPy 创建图像矩阵的方法。
4. 理解普通数组加法和 `cv2.add()` 饱和加法在图像处理中的差异。

## 二、实验原理

数字图像在计算机中通常以矩阵形式存储。灰度图像可以看作二维矩阵，每个元素表示一个像素的灰度值；彩色图像通常由多个通道组成，例如 OpenCV 默认读取的彩色图像通道顺序为 BGR。`cv2.imread()` 可将图像文件读取为 NumPy 数组，`cv2.imwrite()` 可将数组保存为图像文件，`cv2.resize()` 可改变图像尺寸。

创建图像时，可以直接使用 NumPy 构造指定大小和数据类型的数组。例如，元素全部为 0 的 `uint8` 三通道数组对应黑色图像。图像加法本质上是对对应位置的像素值进行相加。普通 NumPy 加法按照 `uint8` 数据类型进行运算，超过 255 时会发生回绕；OpenCV 的 `cv2.add()` 使用饱和运算，像素值超过 255 时会被截断为 255。

## 三、实验结果分析

本实验首先读取花朵图像 `Fig0651(a)(flower_no_compression).tif`，原始图像尺寸为 1600x1200，通道数为 3。使用 `cv2.resize()` 将宽度和高度都缩小为原来的 1/2 后，输出图像尺寸变为 800x600，图像主体内容保持不变，但像素数量减少。

随后使用 `numpy.zeros((50, 50, 3), dtype=np.uint8)` 创建 50x50 黑色正方形图像。由于所有像素值均为 0，显示和保存后的图像为纯黑色。

图像加法部分选取建筑图和器皿图作为输入，并统一缩放为 512x512、三通道图像。使用普通 NumPy 加法得到的结果中，部分像素值超过 255 后发生回绕，因此亮度和颜色会出现不自然变化；使用 `cv2.add()` 得到的结果采用饱和处理，超过 255 的像素值被置为 255，因此整体更偏亮。两种加法结果的平均绝对差异为 65.09，说明二者在像素溢出处理上存在明显差别。

## 四、程序源代码

源码仓库：

```text
{REPO_URL}
```

本实验源码文件：

```text
{SOURCE_FILE}
```

本地源码路径：

```text
{LOCAL_SOURCE_PATH}
```

## 五、程序说明

运行环境：

```bash
pip install -r requirements.txt
```

运行命令：

```bash
python E:\\实验报告\\实验\\src\\exp1_basic_operations.py
```

如果本机 OpenCV 支持 GUI 窗口显示，可使用：

```bash
python E:\\实验报告\\实验\\src\\exp1_basic_operations.py --show
```

程序运行后，结果图像保存到：

```text
E:\\实验报告\\实验\\outputs\\exp1
```

## 六、结果文件说明

- `04_read_write_display_summary.png`：原图读取、1/2 缩放图、50x50 黑色图像对比。
- `10_addition_summary.png`：建筑图、器皿图、NumPy 加法、`cv2.add()` 加法和差异图对比。
- `exp1_result_summary.txt`：实验参数和输出文件摘要。
"""


def set_run_font(run, size: float = 10.5, bold: bool = False) -> None:
    run.font.name = "SimSun"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    run.bold = bold


def add_para(document: Document, text: str, size: float = 10.5):
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    set_run_font(run, size=size)
    return paragraph


def add_heading(document: Document, text: str, level: int = 1):
    paragraph = document.add_heading("", level=level)
    run = paragraph.add_run(text)
    set_run_font(run, size=14 if level == 1 else 12, bold=True)
    return paragraph


def add_code_block(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    run.font.size = Pt(9.5)


def build_document() -> Document:
    document = Document()
    document.styles["Normal"].font.name = "SimSun"
    document.styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    document.styles["Normal"].font.size = Pt(10.5)

    add_heading(document, "实验一  基本图像操作", 1)

    add_heading(document, "一、实验目的", 2)
    add_para(document, "1. 掌握图像处理相关软件和 Python OpenCV 运行环境的基本使用方法。")
    add_para(document, "2. 掌握使用 OpenCV 实现图像读取、写入、显示和缩放操作。")
    add_para(document, "3. 掌握使用 NumPy 创建图像矩阵的方法。")
    add_para(document, "4. 理解普通数组加法和 cv2.add() 饱和加法在图像处理中的差异。")

    add_heading(document, "二、实验原理", 2)
    add_para(
        document,
        "数字图像在计算机中通常以矩阵形式存储。灰度图像可以看作二维矩阵，每个元素表示一个像素的灰度值；彩色图像通常由多个通道组成，例如 OpenCV 默认读取的彩色图像通道顺序为 BGR。cv2.imread() 可将图像文件读取为 NumPy 数组，cv2.imwrite() 可将数组保存为图像文件，cv2.resize() 可改变图像尺寸。",
    )
    add_para(
        document,
        "创建图像时，可以直接使用 NumPy 构造指定大小和数据类型的数组。例如，元素全部为 0 的 uint8 三通道数组对应黑色图像。图像加法本质上是对对应位置的像素值进行相加。普通 NumPy 加法按照 uint8 数据类型进行运算，超过 255 时会发生回绕；OpenCV 的 cv2.add() 使用饱和运算，像素值超过 255 时会被截断为 255。",
    )

    add_heading(document, "三、实验结果分析", 2)
    add_para(
        document,
        "本实验首先读取花朵图像 Fig0651(a)(flower_no_compression).tif，原始图像尺寸为 1600x1200，通道数为 3。使用 cv2.resize() 将宽度和高度都缩小为原来的 1/2 后，输出图像尺寸变为 800x600，图像主体内容保持不变，但像素数量减少。",
    )
    add_para(
        document,
        "随后使用 numpy.zeros((50, 50, 3), dtype=np.uint8) 创建 50x50 黑色正方形图像。由于所有像素值均为 0，显示和保存后的图像为纯黑色。",
    )
    add_para(
        document,
        "图像加法部分选取建筑图和器皿图作为输入，并统一缩放为 512x512、三通道图像。使用普通 NumPy 加法得到的结果中，部分像素值超过 255 后发生回绕，因此亮度和颜色会出现不自然变化；使用 cv2.add() 得到的结果采用饱和处理，超过 255 的像素值被置为 255，因此整体更偏亮。两种加法结果的平均绝对差异为 65.09，说明二者在像素溢出处理上存在明显差别。",
    )

    for title, filename in [
        ("图 1 图像读取、写入、显示与缩放结果", "04_read_write_display_summary.png"),
        ("图 2 图像加法结果对比", "10_addition_summary.png"),
    ]:
        add_para(document, title)
        document.add_picture(str(OUTPUT_DIR / filename), width=Inches(6.4))

    add_heading(document, "四、程序源代码", 2)
    add_para(document, "源码仓库：")
    add_code_block(document, REPO_URL)
    add_para(document, "本实验源码文件：")
    add_code_block(document, SOURCE_FILE)
    add_para(document, "本地源码路径：")
    add_code_block(document, str(LOCAL_SOURCE_PATH))

    add_heading(document, "五、程序说明", 2)
    add_para(document, "运行环境：")
    add_code_block(document, "pip install -r requirements.txt")
    add_para(document, "运行命令：")
    add_code_block(document, r"python E:\实验报告\实验\src\exp1_basic_operations.py")
    add_para(document, "如果本机 OpenCV 支持 GUI 窗口显示，可使用：")
    add_code_block(document, r"python E:\实验报告\实验\src\exp1_basic_operations.py --show")
    add_para(document, "程序运行后，结果图像保存到：")
    add_code_block(document, r"E:\实验报告\实验\outputs\exp1")

    add_heading(document, "六、结果文件说明", 2)
    add_para(document, "04_read_write_display_summary.png：原图读取、1/2 缩放图、50x50 黑色图像对比。")
    add_para(document, "10_addition_summary.png：建筑图、器皿图、NumPy 加法、cv2.add() 加法和差异图对比。")
    add_para(document, "exp1_result_summary.txt：实验参数和输出文件摘要。")

    return document


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    (OUTPUT_DIR / "实验一报告内容.md").write_text(REPORT_MARKDOWN, encoding="utf-8")
    for name in ["exp1_report_content.docx", "exp1_report_content_fixed.docx"]:
        build_document().save(OUTPUT_DIR / name)
        print(OUTPUT_DIR / name)
    print(OUTPUT_DIR / "实验一报告内容.md")


if __name__ == "__main__":
    main()

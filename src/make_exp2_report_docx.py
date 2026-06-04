from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = PROJECT_ROOT / "outputs" / "exp2"


def set_run_font(run, size: float = 10.5, bold: bool = False) -> None:
    run.font.name = "SimSun"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    run.bold = bold


def add_para(document: Document, text: str, size: float = 10.5):
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    set_run_font(run, size=size)
    return paragraph


def add_heading(document: Document, text: str, level: int = 1):
    paragraph = document.add_heading("", level=level)
    run = paragraph.add_run(text)
    set_run_font(run, size=14 if level == 1 else 12, bold=True)
    return paragraph


def build_document() -> Document:
    document = Document()
    document.styles["Normal"].font.name = "SimSun"
    document.styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    document.styles["Normal"].font.size = Pt(10.5)

    add_heading(document, "实验二  图像变换", 1)

    add_heading(document, "一、实验目的", 2)
    add_para(document, "1. 掌握图像几何变换的基本方法。")
    add_para(document, "2. 掌握图像平滑的基本方法。")
    add_para(document, "3. 掌握图像阈值处理的基本方法。")

    add_heading(document, "二、实验原理", 2)
    add_para(
        document,
        "图像几何变换通过坐标映射改变像素在图像中的空间位置。缩放使用插值方法计算新图像像素值，旋转和平移可由仿射变换矩阵实现，透视变换使用 3x3 变换矩阵将原图四点映射到目标图四点。",
    )
    add_para(
        document,
        "图像平滑通过邻域运算降低图像噪声。均值滤波用邻域像素平均值替代中心像素；高斯滤波根据距离中心点的远近赋予不同权重；中值滤波用邻域像素的中位数替代中心像素，通常对椒盐噪声效果较好。",
    )
    add_para(
        document,
        "阈值处理将灰度图像按灰度值分为不同区域。固定阈值需要人工给定阈值，Otsu 算法可根据类间方差自动选取阈值，自适应阈值则根据局部邻域计算阈值，适合处理光照不均匀的图像。",
    )

    add_heading(document, "三、实验结果分析", 2)
    add_para(
        document,
        "几何变换中，缩放 10% 后图像尺寸明显变小，采用 cv2.INTER_AREA 插值可以较好地保留缩小后的整体结构。以图像中心左旋 45 度并保持原尺寸时，部分图像内容会落到边界外，边界处出现黑色填充。平移操作使图像整体向右 10 像素、向下 5 像素移动，移出区域被黑色背景填充。透视变换通过四点对应关系改变图像视角，直线结构仍保持为直线。",
    )
    add_para(
        document,
        "平滑处理中，随着卷积核由 3x3 增大到 9x9，均值滤波和高斯滤波的去噪能力增强，但图像细节和边缘也更模糊。中值滤波对椒盐噪声的抑制更明显，在去除孤立黑白噪点的同时能较好保留边缘轮廓。",
    )
    add_para(
        document,
        "阈值处理中，固定阈值 127 的二值化结果依赖人工阈值选取；Otsu 算法自动计算得到阈值 42，能根据图像灰度分布获得更客观的分割结果；自适应阈值根据局部区域计算阈值，对灰度变化不均匀的图像有更好的局部分割效果。",
    )

    for title, filename in [
        ("图 1 几何变换结果对比", "08_geometry_summary.png"),
        ("图 2 图像平滑结果对比", "12_smoothing_summary.png"),
        ("图 3 阈值处理结果对比", "19_threshold_summary.png"),
    ]:
        add_para(document, title)
        document.add_picture(str(OUTPUT_DIR / filename), width=Inches(6.4))

    add_heading(document, "四、程序源代码", 2)
    add_para(document, r"源代码文件：E:\实验报告\实验\src\exp2_image_transform.py")

    add_heading(document, "五、程序说明", 2)
    add_para(document, "运行环境建议安装 opencv-python 和 numpy。")
    add_para(document, r"运行命令：python E:\实验报告\实验\src\exp2_image_transform.py")
    add_para(
        document,
        r"若本机 OpenCV 支持 GUI 窗口显示，可运行：python E:\实验报告\实验\src\exp2_image_transform.py --show",
    )
    add_para(document, r"程序运行后，结果图像保存到：E:\实验报告\实验\outputs\exp2")

    return document


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    for name in ["exp2_report_content.docx", "exp2_report_content_fixed.docx"]:
        build_document().save(OUTPUT_DIR / name)
        print(OUTPUT_DIR / name)


if __name__ == "__main__":
    main()
